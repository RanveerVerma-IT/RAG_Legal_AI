"""
Document Reader Module
Handles reading and extracting text/data from uploaded documents (PDF, DOCX, etc.)
"""
import re
from io import BytesIO
from typing import Dict, Any, Optional, List
from pathlib import Path

try:
    import pdfplumber
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    from docx import Document
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False


def read_pdf(file_bytes: bytes) -> str:
    """
    Extract text from a PDF file.
    
    Args:
        file_bytes: PDF file content as bytes
        
    Returns:
        Extracted text as string
    """
    if not PDF_SUPPORT:
        raise ImportError("pdfplumber is required for PDF reading. Install it with: pip install pdfplumber")
    
    text_content = []
    try:
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    text_content.append(text)
    except Exception as e:
        raise ValueError(f"Error reading PDF: {str(e)}")
    
    return "\n".join(text_content)


def read_docx(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX file.
    
    Args:
        file_bytes: DOCX file content as bytes
        
    Returns:
        Extracted text as string
    """
    if not DOCX_SUPPORT:
        raise ImportError("python-docx is required for DOCX reading")
    
    try:
        doc = Document(BytesIO(file_bytes))
        paragraphs = [para.text for para in doc.paragraphs]
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        
        return "\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Error reading DOCX: {str(e)}")


def read_document(file_bytes: bytes, file_type: str) -> str:
    """
    Read document based on file type.
    
    Args:
        file_bytes: File content as bytes
        file_type: File MIME type or extension (e.g., 'application/pdf', 'pdf', 'docx')
        
    Returns:
        Extracted text from the document
    """
    file_type_lower = file_type.lower()
    
    if 'pdf' in file_type_lower or file_type_lower.endswith('.pdf'):
        return read_pdf(file_bytes)
    elif 'word' in file_type_lower or 'docx' in file_type_lower or file_type_lower.endswith('.docx'):
        return read_docx(file_bytes)
    elif 'doc' in file_type_lower and 'docx' not in file_type_lower:
        raise ValueError("DOC format is not supported. Please convert to DOCX first.")
    else:
        raise ValueError(f"Unsupported file type: {file_type}. Supported types: PDF, DOCX")


def extract_entities_from_document(text: str) -> Dict[str, Any]:
    """
    Extract common entities (name, DOB, address, etc.) from document text.
    This is a basic extraction - can be enhanced with NLP models.
    
    Args:
        text: Document text content
        
    Returns:
        Dictionary of extracted entities
    """
    entities = {}
    text_lower = text.lower()
    
    # Extract name patterns
    name_patterns = [
        r"name[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
        r"applicant[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
        r"child['\s]*s?\s+name[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
    ]
    for pattern in name_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["applicant_name"] = match.group(1).strip()
            break
    
    # Extract father's name
    father_patterns = [
        r"father['\s]*s?\s+name[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
        r"father[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
    ]
    for pattern in father_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["father_name"] = match.group(1).strip()
            break
    
    # Extract mother's name
    mother_patterns = [
        r"mother['\s]*s?\s+name[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
        r"mother[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
    ]
    for pattern in mother_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["mother_name"] = match.group(1).strip()
            break
    
    # Extract date of birth
    dob_patterns = [
        r"date\s+of\s+birth[:\s]+([0-9]{1,2}[\/\-][0-9]{1,2}[\/\-][0-9]{2,4})",
        r"dob[:\s]+([0-9]{1,2}[\/\-][0-9]{1,2}[\/\-][0-9]{2,4})",
        r"born\s+on[:\s]+([0-9]{1,2}[\/\-][0-9]{1,2}[\/\-][0-9]{2,4})",
    ]
    for pattern in dob_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["date_of_birth"] = match.group(1).strip()
            break
    
    # Extract place of birth
    place_patterns = [
        r"place\s+of\s+birth[:\s]+([A-Z][a-zA-Z\s,]+)",
        r"born\s+at[:\s]+([A-Z][a-zA-Z\s,]+)",
        r"birth\s+place[:\s]+([A-Z][a-zA-Z\s,]+)",
    ]
    for pattern in place_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["place_of_birth"] = match.group(1).strip()
            break
    
    # Extract address
    address_patterns = [
        r"address[:\s]+([A-Z0-9][a-zA-Z0-9\s,.\-\n]+?)(?:\n\n|\n[A-Z]+\s*:|$)",
        r"residential\s+address[:\s]+([A-Z0-9][a-zA-Z0-9\s,.\-\n]+?)(?:\n\n|\n[A-Z]+\s*:|$)",
    ]
    for pattern in address_patterns:
        match = re.search(pattern, text, re.IGNORECASE | re.MULTILINE)
        if match:
            addr = match.group(1).strip()
            # Clean up address (remove extra newlines)
            addr = re.sub(r'\n+', ', ', addr)
            entities["address"] = addr[:300]  # Limit length
            break
    
    # Extract email
    email_match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.IGNORECASE)
    if email_match:
        entities["email"] = email_match.group(0)
    
    # Extract phone
    phone_match = re.search(r"\b(?:\+91[-\s]?)?[6-9][0-9]{9}\b", text)
    if phone_match:
        entities["phone"] = phone_match.group(0)
        entities["contact_phone"] = phone_match.group(0)
    
    # Marriage-specific entity extraction
    # Extract groom's name
    groom_patterns = [
        r"groom['\s]*s?\s+name[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
        r"groom[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
    ]
    for pattern in groom_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["groom_name"] = match.group(1).strip()
            break
    
    # Extract bride's name
    bride_patterns = [
        r"bride['\s]*s?\s+name[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
        r"bride[:\s]+([A-Z][a-zA-Z]+(?:\s+[A-Z][a-zA-Z]+)+)",
    ]
    for pattern in bride_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["bride_name"] = match.group(1).strip()
            break
    
    # Extract marriage date
    marriage_date_patterns = [
        r"marriage\s+date[:\s]+([0-9]{1,2}[\/\-][0-9]{1,2}[\/\-][0-9]{2,4})",
        r"married\s+on[:\s]+([0-9]{1,2}[\/\-][0-9]{1,2}[\/\-][0-9]{2,4})",
        r"wedding\s+date[:\s]+([0-9]{1,2}[\/\-][0-9]{1,2}[\/\-][0-9]{2,4})",
    ]
    for pattern in marriage_date_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            entities["marriage_date"] = match.group(1).strip()
            break
    
    return entities

