"""
Marriage Certificate Application Form Generator
Generates marriage certificate application forms in DOCX and PDF formats
"""
from io import BytesIO
from typing import Dict, Any
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.lib import colors


def _standardize(data: Dict[str, Any]) -> Dict[str, Any]:
    """Standardize and ensure all fields are present"""
    return {
        "groom_name": data.get("groom_name", ""),
        "groom_father_name": data.get("groom_father_name", ""),
        "groom_mother_name": data.get("groom_mother_name", ""),
        "groom_date_of_birth": data.get("groom_date_of_birth", ""),
        "groom_age": data.get("groom_age", ""),
        "groom_address": data.get("groom_address", ""),
        "groom_occupation": data.get("groom_occupation", ""),
        "groom_religion": data.get("groom_religion", ""),
        "bride_name": data.get("bride_name", ""),
        "bride_father_name": data.get("bride_father_name", ""),
        "bride_mother_name": data.get("bride_mother_name", ""),
        "bride_date_of_birth": data.get("bride_date_of_birth", ""),
        "bride_age": data.get("bride_age", ""),
        "bride_address": data.get("bride_address", ""),
        "bride_occupation": data.get("bride_occupation", ""),
        "bride_religion": data.get("bride_religion", ""),
        "marriage_date": data.get("marriage_date", ""),
        "marriage_place": data.get("marriage_place", ""),
        "marriage_venue_address": data.get("marriage_venue_address", ""),
        "witness1_name": data.get("witness1_name", ""),
        "witness1_address": data.get("witness1_address", ""),
        "witness2_name": data.get("witness2_name", ""),
        "witness2_address": data.get("witness2_address", ""),
        "applicant_name": data.get("applicant_name", ""),
        "applicant_relation": data.get("applicant_relation", ""),
        "contact_email": data.get("contact_email", ""),
        "contact_phone": data.get("contact_phone", ""),
    }


def generate_marriage_certificate_docx(data: Dict[str, Any]) -> bytes:
    """Generate marriage certificate application form in DOCX format"""
    values = _standardize(data)
    doc = Document()
    
    # Title
    title = doc.add_heading("MARRIAGE CERTIFICATE APPLICATION FORM", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Header
    para = doc.add_paragraph("Municipal Corporation, Jaipur")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.runs[0].bold = True
    
    para = doc.add_paragraph("Rajasthan, India")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("")  # Spacing
    
    # Groom's Details Section
    doc.add_heading("GROOM'S DETAILS", level=1)
    
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid"
    
    def add_row(label: str, value: str):
        cells = table.add_row().cells
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = value or "—"
    
    add_row("Full Name", values["groom_name"])
    add_row("Father's Name", values["groom_father_name"])
    add_row("Mother's Name", values["groom_mother_name"])
    add_row("Date of Birth", values["groom_date_of_birth"])
    add_row("Age", values["groom_age"])
    add_row("Address", values["groom_address"])
    add_row("Occupation", values["groom_occupation"])
    add_row("Religion", values["groom_religion"])
    
    doc.add_paragraph("")  # Spacing
    
    # Bride's Details Section
    doc.add_heading("BRIDE'S DETAILS", level=1)
    
    table2 = doc.add_table(rows=0, cols=2)
    table2.style = "Light Grid"
    
    add_row("Full Name", values["bride_name"])
    add_row("Father's Name", values["bride_father_name"])
    add_row("Mother's Name", values["bride_mother_name"])
    add_row("Date of Birth", values["bride_date_of_birth"])
    add_row("Age", values["bride_age"])
    add_row("Address", values["bride_address"])
    add_row("Occupation", values["bride_occupation"])
    add_row("Religion", values["bride_religion"])
    
    doc.add_paragraph("")  # Spacing
    
    # Marriage Details Section
    doc.add_heading("MARRIAGE DETAILS", level=1)
    
    table3 = doc.add_table(rows=0, cols=2)
    table3.style = "Light Grid"
    
    add_row("Date of Marriage", values["marriage_date"])
    add_row("Place of Marriage", values["marriage_place"])
    add_row("Marriage Venue Address", values["marriage_venue_address"])
    
    doc.add_paragraph("")  # Spacing
    
    # Witnesses Section
    doc.add_heading("WITNESSES", level=1)
    
    table4 = doc.add_table(rows=0, cols=2)
    table4.style = "Light Grid"
    
    add_row("Witness 1 - Name", values["witness1_name"])
    add_row("Witness 1 - Address", values["witness1_address"])
    add_row("Witness 2 - Name", values["witness2_name"])
    add_row("Witness 2 - Address", values["witness2_address"])
    
    doc.add_paragraph("")  # Spacing
    
    # Applicant Details Section
    doc.add_heading("APPLICANT DETAILS", level=1)
    
    table5 = doc.add_table(rows=0, cols=2)
    table5.style = "Light Grid"
    
    add_row("Applicant Name", values["applicant_name"])
    add_row("Relation to Bride/Groom", values["applicant_relation"])
    add_row("Contact Email", values["contact_email"])
    add_row("Contact Phone", values["contact_phone"])
    
    doc.add_paragraph("")  # Spacing
    
    # Declaration
    para = doc.add_paragraph()
    para.add_run("DECLARATION: ").bold = True
    para.add_run("I hereby declare that the information provided above is true and correct to the best of my knowledge.")
    
    doc.add_paragraph("")  # Spacing
    doc.add_paragraph("")  # Spacing
    
    # Signature line
    para = doc.add_paragraph()
    para.add_run("Signature of Applicant: _________________________")
    para.add_run("\nDate: " + datetime.now().strftime("%d-%m-%Y"))
    
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_marriage_certificate_pdf(data: Dict[str, Any]) -> bytes:
    """Generate marriage certificate application form in PDF format"""
    values = _standardize(data)
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    
    left = 30 * mm
    top = height - 30 * mm
    line_height = 14
    section_gap = 20
    
    # Title
    c.setFont("Helvetica-Bold", 16)
    title = "MARRIAGE CERTIFICATE APPLICATION FORM"
    title_width = c.stringWidth(title, "Helvetica-Bold", 16)
    c.drawString((width - title_width) / 2, top, title)
    
    c.setFont("Helvetica-Bold", 12)
    c.drawString((width - c.stringWidth("Municipal Corporation, Jaipur", "Helvetica-Bold", 12)) / 2, top - 20, "Municipal Corporation, Jaipur")
    c.setFont("Helvetica", 11)
    c.drawString((width - c.stringWidth("Rajasthan, India", "Helvetica", 11)) / 2, top - 35, "Rajasthan, India")
    
    y_pos = top - 60
    
    def draw_field(label: str, value: str):
        nonlocal y_pos
        c.setFont("Helvetica-Bold", 10)
        c.drawString(left, y_pos, f"{label}:")
        c.setFont("Helvetica", 10)
        c.drawString(left + 80, y_pos, value or "—")
        y_pos -= line_height
    
    def draw_section(title: str):
        nonlocal y_pos
        y_pos -= section_gap
        c.setFont("Helvetica-Bold", 12)
        c.drawString(left, y_pos, title)
        c.line(left, y_pos - 2, width - left, y_pos - 2)
        y_pos -= 15
    
    # Groom's Details
    draw_section("GROOM'S DETAILS")
    draw_field("Full Name", values["groom_name"])
    draw_field("Father's Name", values["groom_father_name"])
    draw_field("Mother's Name", values["groom_mother_name"])
    draw_field("Date of Birth", values["groom_date_of_birth"])
    draw_field("Age", values["groom_age"])
    draw_field("Address", values["groom_address"])
    draw_field("Occupation", values["groom_occupation"])
    draw_field("Religion", values["groom_religion"])
    
    # Bride's Details
    draw_section("BRIDE'S DETAILS")
    draw_field("Full Name", values["bride_name"])
    draw_field("Father's Name", values["bride_father_name"])
    draw_field("Mother's Name", values["bride_mother_name"])
    draw_field("Date of Birth", values["bride_date_of_birth"])
    draw_field("Age", values["bride_age"])
    draw_field("Address", values["bride_address"])
    draw_field("Occupation", values["bride_occupation"])
    draw_field("Religion", values["bride_religion"])
    
    # Marriage Details
    draw_section("MARRIAGE DETAILS")
    draw_field("Date of Marriage", values["marriage_date"])
    draw_field("Place of Marriage", values["marriage_place"])
    draw_field("Marriage Venue Address", values["marriage_venue_address"])
    
    # Witnesses
    draw_section("WITNESSES")
    draw_field("Witness 1 - Name", values["witness1_name"])
    draw_field("Witness 1 - Address", values["witness1_address"])
    draw_field("Witness 2 - Name", values["witness2_name"])
    draw_field("Witness 2 - Address", values["witness2_address"])
    
    # Applicant Details
    draw_section("APPLICANT DETAILS")
    draw_field("Applicant Name", values["applicant_name"])
    draw_field("Relation to Bride/Groom", values["applicant_relation"])
    draw_field("Contact Email", values["contact_email"])
    draw_field("Contact Phone", values["contact_phone"])
    
    # Declaration
    y_pos -= section_gap
    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y_pos, "DECLARATION:")
    y_pos -= line_height
    c.setFont("Helvetica", 9)
    c.drawString(left, y_pos, "I hereby declare that the information provided above is true and correct")
    y_pos -= line_height
    c.drawString(left, y_pos, "to the best of my knowledge.")
    
    y_pos -= 20
    c.drawString(left, y_pos, "Signature of Applicant: _________________________")
    y_pos -= line_height
    c.drawString(left, y_pos, f"Date: {datetime.now().strftime('%d-%m-%Y')}")
    
    c.showPage()
    c.save()
    return buf.getvalue()

