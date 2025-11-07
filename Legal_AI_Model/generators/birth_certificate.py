from io import BytesIO
from typing import Dict, Any

from docx import Document
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm


def _standardize(data: Dict[str, Any]) -> Dict[str, Any]:
	# Handle both 'child_name' and 'applicant_name' for backward compatibility
	child_name = data.get("child_name") or data.get("applicant_name", "")
	
	return {
		"applicant_name": child_name,
		"father_name": data.get("father_name", ""),
		"mother_name": data.get("mother_name", ""),
		"address": data.get("address", ""),
		"place_of_birth": data.get("place_of_birth", ""),
		"date_of_birth": data.get("date_of_birth", ""),
		"gender": data.get("gender", ""),
		"email": data.get("email", ""),
		"phone": data.get("phone", ""),
	}


def generate_birth_certificate_docx(data: Dict[str, Any]) -> bytes:
	values = _standardize(data)
	doc = Document()
	title = doc.add_heading("Government of Rajasthan", 0)
	para = doc.add_paragraph("Municipal Corporation / Registrar of Births & Deaths")
	para = doc.add_paragraph("Birth Certificate")
	para.runs[0].bold = True

	table = doc.add_table(rows=0, cols=2)
	table.style = "Light Grid"
	def row(label: str, value: str):
		cells = table.add_row().cells
		cells[0].text = label
		cells[1].text = value or "—"

	row("Name of Child", values["applicant_name"])
	row("Father's Name", values["father_name"])
	row("Mother's Name", values["mother_name"])
	row("Date of Birth", values["date_of_birth"])
	row("Place of Birth", values["place_of_birth"])
	row("Address", values["address"])
	row("Email", values["email"])
	row("Phone", values["phone"])

	doc.add_paragraph("")
	doc.add_paragraph("This is to certify that the above particulars are as per the records.")

	buf = BytesIO()
	doc.save(buf)
	return buf.getvalue()


def generate_birth_certificate_pdf(data: Dict[str, Any]) -> bytes:
	values = _standardize(data)
	buf = BytesIO()
	c = canvas.Canvas(buf, pagesize=A4)
	width, height = A4

	left = 25 * mm
	top = height - 25 * mm

	c.setFont("Helvetica-Bold", 16)
	c.drawString(left, top, "Government of Rajasthan")
	c.setFont("Helvetica", 12)
	c.drawString(left, top - 18, "Municipal Corporation / Registrar of Births & Deaths")
	c.setFont("Helvetica-Bold", 14)
	c.drawString(left, top - 40, "Birth Certificate")

	c.setFont("Helvetica", 11)
	line_y = top - 70
	line_gap = 16

	def line(label: str, value: str):
		nonlocal line_y
		c.drawString(left, line_y, f"{label}: ")
		c.drawString(left + 150, line_y, value or "—")
		line_y -= line_gap

	line("Name of Child", values["applicant_name"])
	line("Father's Name", values["father_name"])
	line("Mother's Name", values["mother_name"])
	line("Date of Birth", values["date_of_birth"])
	line("Place of Birth", values["place_of_birth"])
	line("Address", values["address"])
	line("Email", values["email"])
	line("Phone", values["phone"])

	c.setFont("Helvetica-Oblique", 10)
	c.drawString(left, line_y - 10, "This is to certify that the above particulars are as per the records.")

	c.showPage()
	c.save()
	return buf.getvalue()


